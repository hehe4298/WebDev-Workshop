import os
from datetime import datetime
from docx import Document
from collections import defaultdict

from src.utils.logger import logger
from src.database.db import get_session_context
from src.models.models import Article, Report
from src.config.settings import CATEGORIES

def generate_report() -> str:
    """
    Generates a DOCX report of all currently stored articles grouped by category.
    Returns the path to the generated report or an empty string on failure.
    """
    logger.info("Report generation start")

    reports_dir = "reports"
    if not os.path.exists(reports_dir):
        os.makedirs(reports_dir)

    timestamp = datetime.now()
    filename = f"tech_report_{timestamp.strftime('%Y_%m_%d_%H_%M')}.docx"
    filepath = os.path.join(reports_dir, filename)

    try:
        doc = Document()
        doc.add_heading('Tech Intelligence Report', 0)

        # Add generated timestamp
        doc.add_paragraph(f"Generated at: {timestamp.strftime('%Y-%m-%d %H:%M:%S')}")

        with get_session_context() as session:
            # Group articles by category
            # We could fetch all and group in memory or query per category
            # Since this is a simple report of all articles, let's fetch all
            articles = session.query(Article).order_by(Article.published_date.desc()).all()

            grouped_articles = defaultdict(list)
            for article in articles:
                category = article.category if article.category in CATEGORIES else "Other"
                grouped_articles[category].append(article)

            for category in CATEGORIES:
                if category not in grouped_articles:
                    continue

                cat_articles = grouped_articles[category]
                if not cat_articles:
                    continue

                doc.add_heading(category, level=1)

                for article in cat_articles:
                    doc.add_heading(article.title, level=2)
                    doc.add_paragraph(f"Source: {article.source}")
                    published_str = article.published_date.strftime('%Y-%m-%d %H:%M:%S') if article.published_date else "Unknown"
                    doc.add_paragraph(f"Published Date: {published_str}")
                    doc.add_paragraph(f"URL: {article.url}")

                    # First 500 characters of content
                    content_preview = article.content[:500] + "..." if article.content and len(article.content) > 500 else article.content
                    doc.add_paragraph(content_preview)
                    doc.add_paragraph("-" * 40) # separator

            # Save the document
            doc.save(filepath)

            # Log the report generation in the database
            new_report = Report(report_name=filename, generated_at=timestamp)
            session.add(new_report)
            session.commit()

        logger.info(f"Report generation finish. Saved to {filepath}")
        return filepath

    except Exception as e:
        logger.error(f"Failed to generate report: {e}")
        return ""

if __name__ == "__main__":
    generate_report()
